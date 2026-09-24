"""Summaries and Excel exports."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

GROUPS = {
    "party": ("coalesce(p.name, '(no party)')", "Party"),
    "material": ("coalesce(m.name, '(no material)')", "Material"),
    "vehicle": ("t.vehicle_no", "Vehicle"),
    "day": ("substr(t.closed_at, 1, 10)", "Date"),
}


def summary(conn, date_from: str, date_to: str, group: str = "party") -> list[dict]:
    expr, _ = GROUPS.get(group, GROUPS["party"])
    rows = conn.execute(
        f"SELECT {expr} AS name, COUNT(*) AS trips, SUM(t.net_kg) AS net_kg, "
        "SUM(CASE WHEN t.flags <> '' THEN 1 ELSE 0 END) AS flagged "
        "FROM tickets t LEFT JOIN parties p ON p.id = t.party_id LEFT JOIN materials m ON m.id = t.material_id "
        "WHERE t.status = 'closed' AND substr(t.closed_at,1,10) BETWEEN ? AND ? "
        f"GROUP BY {expr} ORDER BY {'name' if group == 'day' else 'net_kg DESC'}",
        (date_from, date_to))
    return [dict(r) for r in rows]


def day_totals(conn, day: str) -> dict:
    row = conn.execute(
        "SELECT COUNT(*) AS trips, coalesce(SUM(net_kg),0) AS net_kg, "
        "SUM(CASE WHEN flags <> '' THEN 1 ELSE 0 END) AS flagged "
        "FROM tickets WHERE status='closed' AND substr(closed_at,1,10) = ?", (day,)).fetchone()
    open_count = conn.execute("SELECT COUNT(*) FROM tickets WHERE status='open'").fetchone()[0]
    return {"trips": row["trips"], "net_kg": row["net_kg"], "flagged": row["flagged"] or 0, "open": open_count}


def _safe(value):
    """Stop spreadsheet formula injection from text typed into the app."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _sheet(ws, headers: list[str], rows: list[list], widths: list[int] | None = None) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E5B8C")
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append([_safe(v) for v in row])
    ws.freeze_panes = "A2"
    for i, width in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _t(kg) -> float | None:
    return None if kg is None else round(kg / 1000, 3)


def tickets_xlsx(tickets: list[dict], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"
    headers = ["Ticket", "Status", "Vehicle", "Party", "Material", "First weighment", "First kg",
               "Second weighment", "Second kg", "Gross kg", "Tare kg", "Net kg", "Net MT",
               "Flags", "Bill", "Remarks"]
    rows = [[t["ticket_no"], t["status"], t["vehicle_no"], t.get("party_name") or "", t.get("material_name") or "",
             t["first_at"], t["first_kg"], t["second_at"] or "", t["second_kg"], t["gross_kg"], t["tare_kg"],
             t["net_kg"], _t(t["net_kg"]), t["flags"], t.get("bill_no") or "", t["remarks"]] for t in tickets]
    _sheet(ws, headers, rows, [18, 9, 13, 24, 16, 22, 10, 22, 10, 10, 10, 10, 9, 22, 16, 30])
    ws.append([])
    closed = [t for t in tickets if t["status"] == "closed"]
    ws.append(["Total closed", "", "", "", "", "", "", "", "", "", "",
               sum(t["net_kg"] or 0 for t in closed), _t(sum(t["net_kg"] or 0 for t in closed))])
    ws.append([title])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def summary_xlsx(rows: list[dict], group: str, date_from: str, date_to: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    label = GROUPS.get(group, GROUPS["party"])[1]
    _sheet(ws, [label, "Trips", "Net kg", "Net MT", "Flagged"],
           [[r["name"], r["trips"], r["net_kg"], _t(r["net_kg"]), r["flagged"]] for r in rows], [28, 8, 14, 12, 9])
    ws.append([])
    ws.append(["Total", sum(r["trips"] for r in rows), sum(r["net_kg"] or 0 for r in rows),
               _t(sum(r["net_kg"] or 0 for r in rows))])
    ws.append([f"Period {date_from} to {date_to}"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def bill_xlsx(bill: dict, company: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bill"
    ws.append([company["name"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([("Tax invoice" if bill["kind"] == "customer" else "Work statement") + f" {bill['bill_no']}"])
    ws.append([f"Party: {bill['party_name']}  GSTIN: {bill['party_gstin'] or '-'}"])
    ws.append([f"Period: {bill['period_from']} to {bill['period_to']}"])
    ws.append([])
    start = ws.max_row + 1
    ws.append(["Ticket", "Date", "Vehicle", "Material", "Gross kg", "Tare kg", "Net kg", "Net MT"])
    for cell in ws[start]:
        cell.font = Font(bold=True)
    for t in bill["tickets"]:
        ws.append([_safe(t["ticket_no"]), t["closed_at"][:16].replace("T", " "), t["vehicle_no"],
                   _safe(t.get("material_name") or ""), t["gross_kg"], t["tare_kg"], t["net_kg"], _t(t["net_kg"])])
    ws.append([])
    for label, value in (("Trips", bill["trips"]), ("Net MT", _t(bill["total_net_kg"])),
                         ("Rate ₹/MT", bill["rate_per_mt"]), ("Amount ₹", bill["amount"]),
                         (f"CGST {bill['gst_pct'] / 2:g}%", bill["cgst"]), (f"SGST {bill['gst_pct'] / 2:g}%", bill["sgst"]),
                         (f"IGST {bill['gst_pct']:g}%", bill["igst"]), ("Total ₹", bill["total"])):
        ws.append(["", "", "", "", "", label, value])
    for i, width in enumerate([18, 17, 13, 18, 10, 12, 12, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
