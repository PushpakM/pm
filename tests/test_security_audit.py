import base64
import io

from openpyxl import load_workbook

from weighbridge import audit, security
from weighbridge.services import reports
from weighbridge.services.weighing import weigh


def test_password_hashing():
    h = security.hash_password("a long enough passphrase")
    assert h.startswith("scrypt$")
    assert security.verify_password("a long enough passphrase", h)
    assert not security.verify_password("wrong", h)
    assert not security.verify_password("x", "garbage")


def test_password_rules():
    assert security.password_problems("short", 12)
    assert security.password_problems("ramesh-weighbridge", 12, "ramesh")
    assert not security.password_problems("mohada crusher 2026", 12, "ramesh")


def test_totp_rfc6238_vector():
    secret = base64.b32encode(b"12345678901234567890").decode()
    assert security.totp_now(secret, at=59) == "287082"
    assert security.verify_totp(secret, "287082", at=59)
    assert security.verify_totp(secret, "287082", at=59 + 30)  # one step of clock drift allowed
    assert not security.verify_totp(secret, "287082", at=59 + 120)
    assert not security.verify_totp(secret, "abc")


def test_audit_chain_detects_tampering(conn, cfg, operator):
    weigh(conn, cfg, operator, vehicle_no="MH34AB1234", weight_kg=40000)
    weigh(conn, cfg, operator, vehicle_no="MH34AB1234", weight_kg=14000)
    ok, count, bad = audit.verify_chain(conn)
    assert ok and count >= 3 and bad is None
    # Someone with direct database access bypasses the append-only trigger:
    conn.execute("DROP TRIGGER audit_no_update")
    conn.execute("UPDATE audit_log SET details = replace(details, '40000', '45000') WHERE action='ticket.opened'")
    ok, _, bad = audit.verify_chain(conn)
    assert not ok and bad is not None


def test_excel_blocks_formula_injection():
    data = reports.tickets_xlsx([{
        "ticket_no": "WB/2627/000001", "status": "closed", "vehicle_no": "MH34AB1234",
        "party_name": "=HYPERLINK(\"http://evil\")", "material_name": "Boulder", "first_at": "2026-09-24T10:00:00+05:30",
        "first_kg": 40000, "second_at": "2026-09-24T11:00:00+05:30", "second_kg": 14000, "gross_kg": 40000,
        "tare_kg": 14000, "net_kg": 26000, "flags": "", "bill_no": "", "remarks": "+cmd"}], "t")
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["D2"].value.startswith("'=") and ws["P2"].value == "'+cmd"
